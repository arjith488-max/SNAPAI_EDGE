"""SnapAI Edge - Document Upload & Processing API"""

import uuid
import os
import shutil
from fastapi import APIRouter, UploadFile, File, HTTPException, Depends, Form
from sqlalchemy.orm import Session
from typing import Optional

from database.db import get_db, Document

router = APIRouter()

UPLOAD_DIR = os.getenv("UPLOAD_DIR", "./uploads")
MAX_SIZE_MB = int(os.getenv("MAX_UPLOAD_SIZE_MB", 50))
ALLOWED_TYPES = {
    "application/pdf": "pdf",
    "text/plain": "txt",
    "text/markdown": "md",
    "text/csv": "csv",
    "application/json": "json",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "image/png": "png",
    "image/jpeg": "jpg",
    "image/webp": "webp",
}


def extract_text_from_file(path: str, file_type: str) -> tuple[str, int]:
    """Extract text content and page count from uploaded file."""
    text = ""
    pages = 1
    try:
        if file_type == "pdf":
            import fitz  # PyMuPDF
            doc = fitz.open(path)
            pages = len(doc)
            text = "\n\n".join(page.get_text() for page in doc)
            doc.close()
        elif file_type in ("txt", "md", "csv", "json"):
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                text = f.read()
        elif file_type == "docx":
            try:
                import docx
                d = docx.Document(path)
                text = "\n".join(p.text for p in d.paragraphs)
            except ImportError:
                text = "[DOCX support requires python-docx. Install with: pip install python-docx]"
    except Exception as e:
        text = f"[Text extraction failed: {str(e)}]"
    return text, pages


@router.post("/upload")
async def upload_document(
    file: UploadFile = File(...),
    collection: str = Form(default="default"),
    db: Session = Depends(get_db),
):
    # File size check
    content = await file.read()
    size_mb = len(content) / (1024 * 1024)
    if size_mb > MAX_SIZE_MB:
        raise HTTPException(400, f"File too large ({size_mb:.1f} MB). Max allowed: {MAX_SIZE_MB} MB.")

    # File type check
    content_type = file.content_type or "application/octet-stream"
    if content_type not in ALLOWED_TYPES:
        # Try extension fallback
        ext = (file.filename or "").rsplit(".", 1)[-1].lower()
        ext_map = {"pdf": "pdf", "txt": "txt", "md": "md", "csv": "csv",
                   "json": "json", "docx": "docx", "png": "png", "jpg": "jpg", "jpeg": "jpg"}
        if ext not in ext_map:
            raise HTTPException(400, f"Unsupported file type: {content_type}. Allowed: PDF, TXT, MD, CSV, JSON, DOCX, images.")
        file_type = ext_map[ext]
    else:
        file_type = ALLOWED_TYPES[content_type]

    # Safe filename
    doc_id = str(uuid.uuid4())
    safe_name = f"{doc_id}.{file_type}"
    os.makedirs(UPLOAD_DIR, exist_ok=True)
    save_path = os.path.join(UPLOAD_DIR, safe_name)

    with open(save_path, "wb") as f:
        f.write(content)

    # Extract text
    text, pages = extract_text_from_file(save_path, file_type)
    char_count = len(text)

    # Save to DB
    doc = Document(
        id=doc_id,
        filename=safe_name,
        original_name=file.filename or "unnamed",
        file_type=file_type,
        file_size=len(content),
        page_count=pages,
        char_count=char_count,
        collection=collection,
        index_status="pending",
        upload_path=save_path,
    )
    db.add(doc)
    db.commit()

    return {
        "id": doc_id,
        "filename": file.filename,
        "file_type": file_type,
        "size_mb": round(size_mb, 2),
        "pages": pages,
        "chars": char_count,
        "collection": collection,
        "index_status": "pending",
        "message": "Document uploaded. Use /api/rag/index to index for search.",
    }


@router.get("")
async def list_documents(collection: Optional[str] = None, db: Session = Depends(get_db)):
    q = db.query(Document)
    if collection:
        q = q.filter_by(collection=collection)
    docs = q.order_by(Document.created_at.desc()).all()
    return [
        {
            "id": d.id,
            "filename": d.original_name,
            "file_type": d.file_type,
            "size_bytes": d.file_size,
            "pages": d.page_count,
            "chars": d.char_count,
            "chunks": d.chunk_count,
            "collection": d.collection,
            "index_status": d.index_status,
            "created_at": d.created_at,
        }
        for d in docs
    ]


@router.delete("/{doc_id}")
async def delete_document(doc_id: str, db: Session = Depends(get_db)):
    doc = db.query(Document).filter_by(id=doc_id).first()
    if not doc:
        raise HTTPException(404, "Document not found")
    # Remove file
    if os.path.exists(doc.upload_path):
        os.remove(doc.upload_path)
    db.delete(doc)
    db.commit()
    return {"deleted": doc_id}


@router.get("/{doc_id}/preview")
async def preview_document(doc_id: str, db: Session = Depends(get_db)):
    doc = db.query(Document).filter_by(id=doc_id).first()
    if not doc:
        raise HTTPException(404, "Document not found")
    text, _ = extract_text_from_file(doc.upload_path, doc.file_type)
    return {
        "id": doc_id,
        "filename": doc.original_name,
        "text_preview": text[:3000],
        "full_length": len(text),
    }
